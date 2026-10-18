"""Split official MT originals into calendar-year copies; preserve the original workbooks."""
import download_us_crashes as d
import openpyxl,csv,gzip,sys
paths=list((d.ROOT/'MT'/'_source').glob('*.xlsx'))
for path in paths:
    name=next((x for x in ['Crash','Person','Vehicle'] if ('-'+x+'Table') in path.name),None)
    if not name:raise RuntimeError('Unknown workbook '+path.name)
    if '--companions' in sys.argv and name=='Crash':continue
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    sheets=[s for s in wb if s.title!='Section 407']
    if len(sheets)!=1:raise RuntimeError('Unexpected data sheets')
    sheet=sheets[0];iterator=sheet.iter_rows(values_only=True);header=list(next(iterator))
    date_i=header.index('Crash Date');handles={};counts={};files={}
    try:
        for row in iterator:
            if not row or all(x is None for x in row):continue
            value=row[date_i]
            try:year=value.year if hasattr(value,'year') else int(str(value)[:4])
            except (ValueError,TypeError):year='_undated'
            if isinstance(year,int) and not 1900<=year<=2099:year='_undated'
            if year not in handles:
                dest=d.ROOT/'MT'/str(year)/('MDT_'+name+'.csv.gz');dest.parent.mkdir(parents=True,exist_ok=True)
                f=gzip.open(dest.with_name(dest.name+'.split.part'),'wt',encoding='utf-8',newline='',compresslevel=3)
                writer=csv.writer(f);writer.writerow(header);handles[year]=(f,writer);counts[year]=0;files[year]=dest
            handles[year][1].writerow([x.isoformat() if hasattr(x,'isoformat') else d.cell(x) for x in row]);counts[year]+=1
    finally:
        for f,writer in handles.values():f.close()
        wb.close()
    for year,dest in files.items():
        temp=dest.with_name(dest.name+'.split.part');n,columns,bad=d.verify_csv(temp)
        if n!=counts[year] or bad:raise RuntimeError('Split CSV validation failed')
        temp.replace(dest)
        receipt={'status':'complete','kind':'year_split_original_excel','region':'MT','year':year,'name':'MDT_'+name,
                 'records':n,'columns':columns,'source_original':d.relative(path),'year_attribution':'Published Crash Date in original table',
                 'coordinate_crs':'projected X/Y; official related map EPSG:32100, confirm more controls before use',
                 'access_date':d.CLIENT_DATE,'path':d.relative(dest),'stored_bytes':dest.stat().st_size,'sha256':d.hash_file(dest)}
        d.atomic_json(dest.with_name(dest.name+'.metadata.json'),receipt)
    print(name,'rows',sum(counts.values()),'year counts',counts,flush=True)
print('MT SPLIT COMPLETE',flush=True)
