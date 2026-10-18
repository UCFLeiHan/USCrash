from __future__ import annotations
import os as _uscrash_paths_os

import argparse
import csv
import hashlib
import json
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import date, datetime, time
from pathlib import Path

ROOT = Path(_uscrash_paths_os.environ['USCRASH_ROOT'])
STATES = ROOT / "States"
MANIFEST_DIR = ROOT / "Unified_MMUCC6" / "manifests"
LOG_DIR = ROOT / "Unified_MMUCC6" / "logs"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(8 * 1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean_headers(values):
    seen = {}
    out = []
    for i, value in enumerate(values, 1):
        name = str(value).strip() if value is not None else ""
        name = name or f"unnamed_{i}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        out.append(name if count == 1 else f"{name}_{count}")
    return out


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return value


def workbook_to_csvs(workbook_path: Path, output_dir: Path, year: int, prefix: str = "MS"):
    from openpyxl import load_workbook

    output_dir.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    records = []
    try:
        for ws in wb.worksheets:
            safe_sheet = re.sub(r"[^A-Za-z0-9]+", "_", ws.title).strip("_") or "Sheet"
            out = output_dir / f"{prefix}_{safe_sheet}_{year}.csv"
            rows = ws.iter_rows(values_only=True)
            try:
                first = next(rows)
            except StopIteration:
                continue
            headers = clean_headers(first)
            row_count = 0
            with out.open("w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f, lineterminator="\n")
                writer.writerow(headers)
                for row in rows:
                    writer.writerow([cell_text(v) for v in row])
                    row_count += 1
            records.append({
                "state": prefix,
                "year": year,
                "source_file": str(workbook_path),
                "source_member_or_layer": ws.title,
                "output_file": str(out),
                "rows": row_count,
                "columns": len(headers),
                "output_sha256": sha256(out),
                "status": "created",
            })
    finally:
        wb.close()
    return records


def convert_ms(start_year=2004, end_year=2025):
    records = []
    for year in range(start_year, end_year + 1):
        year_dir = STATES / "MS" / str(year)
        out_dir = year_dir / "CSV"
        xlsx = year_dir / f"mississippi_ecrash_{year}.xlsx"
        zpath = year_dir / f"mississippi_reportbeam_{year}.zip"
        if xlsx.exists():
            source_hash = sha256(xlsx)
            rs = workbook_to_csvs(xlsx, out_dir, year)
            for r in rs:
                r["source_sha256"] = source_hash
            records.extend(rs)
        elif zpath.exists():
            source_hash = sha256(zpath)
            with zipfile.ZipFile(zpath) as zf, tempfile.TemporaryDirectory(prefix=f"ms_{year}_") as td:
                for member in zf.namelist():
                    if not member.lower().endswith((".xlsx", ".xlsm")):
                        continue
                    extracted = Path(zf.extract(member, td))
                    rs = workbook_to_csvs(extracted, out_dir, year)
                    for r in rs:
                        r["source_file"] = str(zpath)
                        r["source_member_or_layer"] = f"{member}::{r['source_member_or_layer']}"
                        r["source_sha256"] = source_hash
                    records.extend(rs)
        print(f"MS {year}: {len([r for r in records if r['year'] == year])} CSV files", flush=True)
    return records


def copy_oh():
    records = []
    for year in range(2017, 2021):
        src = STATES / "OH" / str(year) / "Raw" / "Figshare_29437694" / f"{year}_clean.csv"
        out_dir = STATES / "OH" / str(year) / "CSV"
        out_dir.mkdir(parents=True, exist_ok=True)
        dst = out_dir / f"OH_{year}_crash_vehicle_person_flat.csv"
        shutil.copy2(src, dst)
        with dst.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
            reader = csv.reader(f)
            header = next(reader)
            rows = sum(1 for _ in reader)
        records.append({
            "state": "OH", "year": year, "source_file": str(src),
            "source_member_or_layer": "", "output_file": str(dst),
            "rows": rows, "columns": len(header), "source_sha256": sha256(src),
            "output_sha256": sha256(dst), "status": "copied",
        })
        print(f"OH {year}: {rows} rows", flush=True)
    return records


def convert_or():
    import fiona
    from pyproj import CRS, Transformer

    records = []
    for year in range(2011, 2015):
        src = STATES / "OR" / str(year) / "Raw" / f"crashes{year}.gdb.zip"
        out_dir = STATES / "OR" / str(year) / "CSV"
        out_dir.mkdir(parents=True, exist_ok=True)
        dst = out_dir / f"OR_{year}_crash.csv"
        vsi = "/vsizip/" + src.as_posix()
        layers = fiona.listlayers(vsi)
        if not layers:
            raise RuntimeError(f"No layer found in {src}")
        layer = next((x for x in layers if "crash" in x.lower()), layers[0])
        with fiona.open(vsi, layer=layer) as collection:
            props = list(collection.schema["properties"].keys())
            src_crs = CRS.from_user_input(collection.crs_wkt or collection.crs)
            transformer = Transformer.from_crs(src_crs, "EPSG:4326", always_xy=True)
            rows = 0
            valid_points = 0
            with dst.open("w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=props + ["longitude_wgs84", "latitude_wgs84"], lineterminator="\n")
                writer.writeheader()
                for feature in collection:
                    row = dict(feature["properties"])
                    lon = lat = ""
                    geom = feature.get("geometry")
                    if geom and geom.get("coordinates"):
                        coords = geom["coordinates"]
                        if geom.get("type") == "Point":
                            lon, lat = transformer.transform(coords[0], coords[1])
                            if -125.0 <= lon <= -116.0 and 41.0 <= lat <= 47.0:
                                valid_points += 1
                    row["longitude_wgs84"] = lon
                    row["latitude_wgs84"] = lat
                    writer.writerow(row)
                    rows += 1
        records.append({
            "state": "OR", "year": year, "source_file": str(src),
            "source_member_or_layer": layer, "output_file": str(dst),
            "rows": rows, "columns": len(props) + 2, "valid_wgs84_points": valid_points,
            "source_sha256": sha256(src), "output_sha256": sha256(dst), "status": "created",
        })
        print(f"OR {year}: {rows} rows; {valid_points} coordinates in Oregon bounds", flush=True)
    return records


def write_manifest(records, label):
    MANIFEST_DIR.mkdir(parents=True, exist_ok=True)
    out_json = MANIFEST_DIR / f"source_format_conversion_{label}.json"
    out_csv = MANIFEST_DIR / f"source_format_conversion_{label}.csv"
    out_json.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    fields = sorted({k for r in records for k in r})
    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(records)
    return out_csv


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--state", choices=["MS", "OH", "OR", "ALL"], default="ALL")
    parser.add_argument("--start-year", type=int, default=2004)
    parser.add_argument("--end-year", type=int, default=2025)
    args = parser.parse_args()
    jobs = {"MS": convert_ms, "OH": copy_oh, "OR": convert_or}
    selected = list(jobs) if args.state == "ALL" else [args.state]
    for state in selected:
        records = convert_ms(args.start_year, args.end_year) if state == "MS" else jobs[state]()
        label = f"MS_{args.start_year}_{args.end_year}" if state == "MS" else state
        manifest = write_manifest(records, label)
        print(f"Manifest: {manifest}", flush=True)


if __name__ == "__main__":
    main()
