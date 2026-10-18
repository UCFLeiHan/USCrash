from __future__ import annotations

import csv
import hashlib
import json
import os
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

DATA_ROOT = Path(os.environ.get("USCRASH_ROOT", r"F:\LH_Research\USCrash\Crash"))
PIPE = DATA_ROOT / "Unified_MMUCC6"
STATES = DATA_ROOT / "States"
PREFLIGHT = PIPE / "qa" / "preflight" / "state_year_preflight.csv"
CANDIDATES = PIPE / "config" / "candidates"
FIELDS = ["crash_id", "crash_datetime", "crash_date", "crash_time", "latitude", "longitude", "severity", "manner_collision"]


def normalized_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.strip().lower()).strip("_")


def read_csv_header(path: Path):
    encodings = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
    for encoding in encodings:
        try:
            with path.open("r", encoding=encoding, errors="strict", newline="") as f:
                return next(csv.reader(f)), encoding
        except (UnicodeDecodeError, StopIteration, csv.Error):
            continue
    raise ValueError("CSV header unreadable")


def read_xlsx_header(path: Path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        values = next(ws.iter_rows(values_only=True))
        return ["" if x is None else str(x) for x in values], f"xlsx:{ws.title}"
    finally:
        wb.close()


def read_header(path: Path):
    if path.suffix.lower() == ".csv": return read_csv_header(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}: return read_xlsx_header(path)
    raise ValueError(f"unsupported extension {path.suffix}")


def score(mapping: dict):
    parts = {
        "identifier": 20 if mapping.get("crash_id") else 0,
        "time": 20 if mapping.get("crash_datetime") or mapping.get("crash_date") else (10 if mapping.get("crash_time") else 0),
        "coordinates": 20 if mapping.get("latitude") and mapping.get("longitude") else 0,
        "severity": 20 if mapping.get("severity") else 0,
        "manner": 20 if mapping.get("manner_collision") else 0,
    }
    return sum(parts.values()), parts


def main():
    CANDIDATES.mkdir(parents=True, exist_ok=True)
    with PREFLIGHT.open(encoding="utf-8-sig", newline="") as f:
        source_rows = list(csv.DictReader(f))
    groups = defaultdict(lambda: defaultdict(list))
    errors = []
    for row in source_rows:
        if not row["year"] or not row["primary_file"]: continue
        path = STATES / Path(row["primary_file"])
        try:
            headers, encoding = read_header(path)
            norm = [normalized_header(x) for x in headers]
            signature = hashlib.sha256("\x1f".join(norm).encode()).hexdigest()[:16]
            groups[row["state"]][signature].append((row, path, headers, encoding))
        except Exception as exc:
            errors.append({"state": row["state"], "year": row["year"], "path": str(path), "error": str(exc)})

    discovery = []
    for state, signatures in sorted(groups.items()):
        candidate = {"state": state, "generated_utc": datetime.now(timezone.utc).isoformat(), "status": "CANDIDATE_NOT_REVIEWED", "schemas": []}
        for index, (signature, items) in enumerate(sorted(signatures.items(), key=lambda x: min(int(i[0]["year"]) for i in x[1])), 1):
            years = sorted(int(item[0]["year"]) for item in items)
            first = items[0]
            mapping = {name: first[0].get(name, "") for name in FIELDS if first[0].get(name, "")}
            total, components = score(mapping)
            schema_id = f"{state}_candidate_{years[0]}_{years[-1]}_{index}"
            schema = {
                "schema_id": schema_id, "header_signature": signature, "years": years,
                "example_source": str(first[1].relative_to(STATES)).replace("\\", "/"),
                "encoding_or_sheet": first[3], "column_count": len(first[2]),
                "field_candidates": mapping, "completeness_score": total,
                "score_components": components, "review_status": "NOT_REVIEWED",
            }
            candidate["schemas"].append(schema)
            discovery.append({
                "state": state, "schema_id": schema_id, "year_start": years[0], "year_end": years[-1],
                "year_count": len(years), "header_signature": signature, "column_count": len(first[2]),
                "completeness_score": total, **components, "example_source": schema["example_source"],
                "review_status": "NOT_REVIEWED",
            })
        (CANDIDATES / f"{state}.candidate.json").write_text(json.dumps(candidate, ensure_ascii=False, indent=2), encoding="utf-8")

    discovery_path = PIPE / "qa" / "preflight" / "schema_discovery.csv"
    with discovery_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(discovery[0]), lineterminator="\n")
        writer.writeheader(); writer.writerows(discovery)
    queue = sorted(discovery, key=lambda r: (-int(r["completeness_score"]), int(r["year_count"]) * -1, r["state"], int(r["year_start"])))
    queue_path = PIPE / "qa" / "preflight" / "manual_review_queue.csv"
    with queue_path.open("w", encoding="utf-8-sig", newline="") as f:
        fields = ["priority"] + list(queue[0])
        writer = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for i, row in enumerate(queue, 1): writer.writerow({"priority": i, **row})
    error_path = PIPE / "qa" / "preflight" / "schema_discovery_errors.csv"
    with error_path.open("w", encoding="utf-8-sig", newline="") as f:
        fields = ["state", "year", "path", "error"]
        writer = csv.DictWriter(f, fieldnames=fields, lineterminator="\n"); writer.writeheader(); writer.writerows(errors)
    result = {"states": len(groups), "schema_versions": len(discovery), "source_years_grouped": sum(int(x["year_count"]) for x in discovery),
              "errors": len(errors), "generated_utc": datetime.now(timezone.utc).isoformat(), "result": "PASS" if not errors else "WARN"}
    (PIPE / "qa" / "preflight" / "schema_discovery_manifest.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
