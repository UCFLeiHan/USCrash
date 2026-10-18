import os as _uscrash_paths_os
import csv
from pathlib import Path
from openpyxl import load_workbook

ROOT = (Path(_uscrash_paths_os.environ['USCRASH_ROOT']) / 'Unified_MMUCC6')
SOURCE = ROOT / "reference" / "MMUCC6_schema_package" / "MMUCC6_schema_package" / "MMUCC6_MappingTemplate.xlsx"
OUTPUT = ROOT / "reference" / "MMUCC6_element_dictionary.csv"

wb = load_workbook(SOURCE, read_only=True, data_only=True)
fields = [
    "level", "class", "identifier", "data_element", "subfield_or_group",
    "definition", "data_type", "cardinality", "allow_nulls", "namespace",
    "property_name", "property_type", "base_type", "contained_in",
    "sub_group_head", "category", "category_definition", "category_cardinality",
    "sample_data", "comments", "exchange_path",
]
rows = []
for ws in wb.worksheets:
    iterator = ws.iter_rows(values_only=True)
    next(iterator)
    parent_identifier = None
    parent_element = None
    for values in iterator:
        values = list(values) + [None] * (20 - len(values))
        if values[1]:
            parent_identifier = values[1]
        if values[2]:
            parent_element = values[2]
        if not any(v is not None for v in values):
            continue
        rows.append(dict(zip(fields, [ws.title] + values[:20])) | {
            "parent_identifier": parent_identifier,
            "parent_data_element": parent_element,
        })
wb.close()

out_fields = fields + ["parent_identifier", "parent_data_element"]
with OUTPUT.open("w", encoding="utf-8-sig", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=out_fields, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
print(f"{OUTPUT}: {len(rows)} rows")
